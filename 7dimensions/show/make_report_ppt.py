#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 7dimensions 测试结果生成 PPT(pptx) + PDF 报告。

数据来源：测试结果目录下每个「图片名文件夹」里的 solutions.json，
其中包含 input_image（测试原图）、evidence_by_dim（每维检索到的 poor 图 / 放进
prompt 的 good 图 / dim_original_text）以及 solutions（每个 direction 方案的 text）。

版式（16:9，每个 direction 方案 4 页）：
  每页左侧    ：测试原图
  第 1 页右侧 ：维度1 Ratio + 维度2 Composition（各含 poor图、good图、
                dim_original_text、拍照缺陷分析、针对性整改建议）
  第 2 页右侧 ：维度3 Camera + 维度4 Position
  第 3 页右侧 ：维度5 Pose  + 维度6 Focus
  第 4 页右侧 ：上方 = 维度7 Color；下方 = 该方案完整的“针对性整改建议”

用法（整批 -> 单份报告，保存到输入测试路径）：
  python make_report_ppt.py \
      --results-dir /workspace/ai-ddge/7dimensions/output/0910_sv_FIRST_GOOD_PER_POOR_qwen38flash1 \
      --save-dir   /workspace/ai-ddge/7dimensions/output/0910_sv_FIRST_GOOD_PER_POOR_qwen38flash1 \
      --data-root  /workspace/ai-ddge/7dimensions/dataset
  # 可选：--only "669da389b6114a208175994fa13af3fe" 只处理某一张
  #       --max-solutions 3 每张只取前 N 个方案
  #       --out-name "0830报告" 自定义输出文件名（默认=<结果目录名>_报告）
说明：
  - 整个结果目录的所有图片会合成「单份」pptx + pdf（不是每张一个）。
  - 每个维度展示 json 中全部 poor/good 图与 dim_original_text；证据多时自动
    缩小图片并按网格排布，保证全部放下（同一 poor 对应多张 good 也会逐一展示）。
"""
from __future__ import annotations

import argparse
import glob
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile

from PIL import Image

# ---------------------------------------------------------------------------
# 依赖检查
# ---------------------------------------------------------------------------
_HAS_PPTX = True
_HAS_RL = True
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
except Exception:  # pragma: no cover
    _HAS_PPTX = False

try:
    from reportlab.pdfgen import canvas as _rl_canvas
    from reportlab.lib.pagesizes import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.utils import ImageReader
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
except Exception as e:  # pragma: no cover
    print("[warn] reportlab 不可用，将无法导出 PDF:", e)
    _HAS_RL = False

# ---------------------------------------------------------------------------
# 布局常量（单位：英寸；16:9）
# ---------------------------------------------------------------------------
PAGE_W = 13.333
PAGE_H = 7.5
MARGIN = 0.25
TITLE_Y = 0.10
TITLE_H = 0.55
CONTENT_Y = 0.85
CONTENT_H = PAGE_H - CONTENT_Y - 0.30          # 6.35
LEFT_X = MARGIN
LEFT_W = 4.0
RIGHT_X = MARGIN + LEFT_W + 0.30               # 4.55
RIGHT_W = PAGE_W - MARGIN - RIGHT_X            # ~8.53
BLOCK_GAP = 0.12
HEAD_H = 0.26
CAP_H = 0.14
PAGE4_TOP_H = 3.3
TEXT_BASE = 9.0
TEXT_MIN = 6.5

BLACK = (0.0, 0.0, 0.0)
GRAY = (0.42, 0.42, 0.42)
DARKBLUE = (0.12, 0.25, 0.42)

DIM_ORDER = ["ratio", "composition", "camera", "position", "pose", "focus", "color"]
LABEL = {
    "ratio": "1 · Ratio 画面比例",
    "composition": "2 · Composition 构图取景",
    "camera": "3 · Camera 机位视角",
    "position": "4 · Position 主体位置",
    "pose": "5 · Pose 姿态动作",
    "focus": "6 · Focus 对焦景深",
    "color": "7 · Color 色彩光影",
}
CN = {
    "ratio": "画面比例", "composition": "构图取景", "camera": "机位视角",
    "position": "主体位置", "pose": "姿态动作", "focus": "对焦景深",
    "color": "色彩光影",
}
PAGES = [["ratio", "composition"], ["camera", "position"],
         ["pose", "focus"], ["color"]]
EVIDENCE_PER_DIM = 0          # 0 = 展示该维全部证据（所有 poor/good 对）；>0 只取前 N 条
DEFAULT_DATA_ROOTS = ["/workspace/ai-ddge/7dimensions/dataset"]

# 多证据网格布局参数（英寸）：优先保证图片尽量大，文字适当缩小让位给图片
PAIR_GAP = 0.12      # 同一条证据内 poor 与 good 的间距
COL_GAP = 0.12       # 网格列间距
ROW_GAP = 0.12       # 网格行间距
PAD = 0.10           # 维度块内边距
# 图片盒尺寸递退档（宽, 高）：从大到小，证据多时逐档缩小；首档明显放大提升可读性
IMG_STEPS = [(1.35, 1.10), (1.20, 1.00), (1.05, 0.88), (0.90, 0.75), (0.75, 0.62)]
FONT_STEPS = [8.0, 7.0, 6.0, 5.5]   # 文字适当缩小，把更多高度让给图片

# ---------------------------------------------------------------------------
# 解析 solutions.json
# ---------------------------------------------------------------------------
_DIM_RE = (r"^(画面比例|构图取景|机位视角|主体位置|姿态动作|姿态|对焦景深|对焦|色彩光影|色彩)"
           r"\s*(?:整改)?\s*[：:]\s*(.*)$")
_DIM_NAME2KEY = {
    "画面比例": "ratio", "构图取景": "composition", "机位视角": "camera",
    "主体位置": "position", "姿态动作": "pose", "姿态": "pose",
    "对焦景深": "focus", "对焦": "focus", "色彩光影": "color", "色彩": "color",
}


def _split_sections(text: str):
    """拆出「拍照缺陷分析」和「针对性整改建议」两个段落。"""
    m_def = text.find("# 拍照缺陷分析")
    m_adv = text.find("# 针对性整改建议")
    defect, advice = "", ""
    if m_def != -1 and (m_adv > m_def or m_adv == -1):
        defect = text[m_def + len("# 拍照缺陷分析"): (m_adv if m_adv > m_def else len(text))]
    if m_adv != -1:
        advice = text[m_adv + len("# 针对性整改建议"):]
    return defect, advice


def _parse_bullets(section: str) -> dict:
    """把段落里的 '- 画面比例：xxx' 类条目解析成 {dim_key: 文本}。"""
    out = {}
    if not section:
        return out
    for line in section.splitlines():
        line = line.strip()
        if not line.startswith("-"):
            continue
        body = line[1:].strip()
        body = re.sub(r"^[（(][^）)]*[）)]\s*", "", body)   # 去掉“（仅使用本组证据）”
        m = re.match(_DIM_RE, body)
        if m:
            key = _DIM_NAME2KEY.get(m.group(1))
            if key:
                txt = m.group(2).strip()
                # 剥掉冒号后紧跟的“（仅使用本组证据）”等括号说明
                txt = re.sub(r"^[（(][^）)]*[）)]\s*", "", txt)
                out[key] = txt
    return out


def parse_solution_text(text: str) -> dict:
    defect, advice = _split_sections(text)
    return {"defect": _parse_bullets(defect), "advice": _parse_bullets(advice)}


# ---------------------------------------------------------------------------
# 图片解析（多根路径回退）
# ---------------------------------------------------------------------------
def resolve_image(path: str, roots: list) -> str:
    """按多种策略解析图片真实路径，找不到返回原路径（渲染时显示“图片缺失”）。"""
    if not path:
        return path
    cands = [path]
    for root in roots:
        if root:
            cands.append(os.path.join(root, path.lstrip("/")))
    cands.append(path.replace("/workspace/ai-ddge/AesRecon",
                              "/workspace/ai-ddge/7dimensions/dataset/AesRecon"))
    for c in cands:
        if c and os.path.exists(c):
            return c
    return path


# ---------------------------------------------------------------------------
# 文本排版工具
# ---------------------------------------------------------------------------
_FULLWIDTH_RANGES = ((0x3000, 0x303F), (0xFF00, 0xFFEF))
_FULLWIDTH_SINGLE = {0x00B7, 0x2018, 0x2019, 0x201C, 0x201D, 0x2013, 0x2014, 0x2026}


def _char_w(ch: str, pt: float) -> float:
    """近似字符宽度（pt）：CJK/全角标点=1.0 字号，ASCII≈0.5 字号。"""
    o = ord(ch)
    if o > 0x2E7F:
        return pt
    for a, b in _FULLWIDTH_RANGES:
        if a <= o <= b:
            return pt
    if o in _FULLWIDTH_SINGLE:
        return pt
    return pt * 0.50


def wrap_text(text: str, width_in: float, pt: float) -> list:
    width_pt = max(width_in * 72.0, 1.0)
    lines, cur, cur_w = [], "", 0.0
    for ch in text:
        if ch == "\n":
            lines.append(cur)
            cur, cur_w = "", 0.0
            continue
        w = _char_w(ch, pt)
        if cur and cur_w + w > width_pt:
            lines.append(cur)
            cur, cur_w = ch, w
        else:
            cur += ch
            cur_w += w
    if cur:
        lines.append(cur)
    return lines


def needed_height(text: str, width_in: float, pt: float) -> float:
    return len(wrap_text(text, width_in, pt)) * pt * 1.25 / 72.0


def fit_pt(text: str, width_in: float, height_in: float,
           base: float = TEXT_BASE, min_pt: float = TEXT_MIN) -> float:
    h0 = needed_height(text, width_in, base)
    if h0 <= height_in:
        return base
    pt = max(min_pt, base * height_in / max(h0, 1e-6))
    while pt > min_pt and needed_height(text, width_in, pt) > height_in:
        pt -= 0.5
    return max(min_pt, pt)


def fit_box(iw, ih, bw, bh):
    if iw <= 0 or ih <= 0:
        return (bw, bh)
    r = min(bw / iw, bh / ih)
    return (iw * r, ih * r)


# ---------------------------------------------------------------------------
# 单页元素构建
# ---------------------------------------------------------------------------
def _dim_text(ev_list, sol, dim):
    """组装一个维度的文本：全部证据原文 + 拍照缺陷分析 + 针对性整改建议。"""
    parts = []
    if not ev_list:
        parts.append("【该维度无检索证据】")
    else:
        for i, ev in enumerate(ev_list, 1):
            parts.append(f"【证据 {i}】{ev['text']}")
    parts.append(f"【拍照缺陷分析】{sol['defect'].get(dim, '')}")
    parts.append(f"【针对性整改建议】{sol['advice'].get(dim, '')}")
    return "\n".join(p for p in parts if p and p.split("】", 1)[-1].strip())


def _img_fit(path, bw, bh, dims_cache):
    """把图片按比例放入 (bw, bh) 盒子，返回实际 (iw, ih)；尺寸按原图缓存。"""
    if not path or not os.path.exists(path):
        return bw, bh
    key = (path, "size")
    if key not in dims_cache:
        try:
            with Image.open(path) as im:
                dims_cache[key] = im.size
        except Exception:
            dims_cache[key] = (1, 1)
    iw, ih = dims_cache[key]
    return fit_box(iw, ih, bw, bh)


def _dim_metrics(ev_list, sol, dim, w, img_w, img_h, pt):
    """给定图盒尺寸/字号，返回 (n_cols, n_rows, grid_h, text_h, text)。"""
    n = len(ev_list)
    text = _dim_text(ev_list, sol, dim)
    if n == 0:
        return 1, 1, 0.0, needed_height(text, w, pt), text
    pair_w = 2 * img_w + PAIR_GAP
    pitch = pair_w + COL_GAP
    n_cols = max(1, int((w - pair_w) / pitch) + 1)
    n_cols = max(1, min(n_cols, n))
    n_rows = math.ceil(n / n_cols)
    grid_h = n_rows * (img_h + CAP_H) + (n_rows - 1) * ROW_GAP
    text_h = needed_height(text, w, pt)
    return n_cols, n_rows, grid_h, text_h, text


def _fit_dim(ev_list, sol, dim, w, h):
    """在高度 h 内拟合一个维度块；返回 (img_w, img_h, pt, metrics)。"""
    for img_w, img_h in IMG_STEPS:
        for pt in FONT_STEPS:
            m = _dim_metrics(ev_list, sol, dim, w, img_w, img_h, pt)
            if HEAD_H + m[2] + m[3] + PAD <= h:
                return img_w, img_h, pt, m
    img_w, img_h = IMG_STEPS[-1]
    pt = FONT_STEPS[-1]
    m = _dim_metrics(ev_list, sol, dim, w, img_w, img_h, pt)
    return img_w, img_h, pt, m


def _dim_block_items(x, y, w, h, ev_list, label, sol, dim, roots, dims_cache):
    """一个维度块：标题 + 全部 poor/good 图网格 + 图注 + 证据/缺陷/整改文本。"""
    items = []
    items.append({"t": "text", "x": x, "y": y, "w": w, "h": HEAD_H,
                  "text": f"【{label}】", "size": 11, "bold": True, "color": DARKBLUE})

    img_w, img_h, pt, (n_cols, n_rows, grid_h, text_h, text) = \
        _fit_dim(ev_list, sol, dim, w, h)

    if ev_list:
        pair_w = 2 * img_w + PAIR_GAP
        pitch = pair_w + COL_GAP
        row_w = (n_cols - 1) * pitch + pair_w
        x0 = x + (w - row_w) / 2
        for i, ev in enumerate(ev_list):
            r, c = divmod(i, n_cols)
            cy = y + HEAD_H + r * (img_h + CAP_H + ROW_GAP)
            cx = x0 + c * pitch
            for side, p in (("poor", ev["poor"]), ("good", ev["good"])):
                iw, ih = _img_fit(p, img_w, img_h, dims_cache)
                ix = cx + ((img_w - iw) / 2 if side == "poor"
                           else img_w + PAIR_GAP + (img_w - iw) / 2)
                iy = cy + (img_h - ih) / 2
                items.append({"t": "image", "x": ix, "y": iy, "w": iw, "h": ih, "path": p})
                items.append({"t": "text", "x": ix, "y": cy + img_h + 0.01,
                              "w": img_w, "h": CAP_H,
                              "text": f"poor{i}" if side == "poor" else f"good{i}",
                              "size": 7, "bold": False, "color": GRAY})

    ty = y + HEAD_H + grid_h + 0.02
    items.append({"t": "text", "x": x, "y": ty, "w": w,
                  "h": max(h - (ty - y), 0.05),
                  "text": text, "size": pt, "bold": False, "color": BLACK})
    return items


def _dim_natural(ev_list, sol, dim, w):
    """基准（最大图/最大字号）下该维的自然块高。"""
    m = _dim_metrics(ev_list, sol, dim, w, IMG_STEPS[0][0], IMG_STEPS[0][1], FONT_STEPS[0])
    return HEAD_H + m[2] + m[3] + PAD


def _dim_min_height(ev_list, sol, dim, w):
    """最小（最小图/最小字号）设置下该维需要的最小块高。"""
    m = _dim_metrics(ev_list, sol, dim, w, IMG_STEPS[-1][0], IMG_STEPS[-1][1], FONT_STEPS[-1])
    return HEAD_H + m[2] + m[3] + PAD


def _page_groups(sol, w):
    """返回实际页面分组：('dims',[dim,...]) 或 ('color', dim)。
    同一页两个维度在最小设置下也放不下、或某维证据过多(>5 条)时，
    拆成两页各占一页（图片更大更清楚）。"""
    groups = []
    free = CONTENT_H - BLOCK_GAP
    for dims in PAGES[:3]:
        if len(dims) == 2:
            nat = [_dim_natural(sol["ev"].get(d) or [], sol, d, w) for d in dims]
            shares = ([free * n / sum(nat) for n in nat] if sum(nat) > free else nat)
            ev_n = [len(sol["ev"].get(d) or []) for d in dims]
            need_split = any(
                _dim_min_height(sol["ev"].get(d) or [], sol, d, w) > s
                for d, s in zip(dims, shares)) or any(n > 5 for n in ev_n)
            if need_split:
                for d in dims:
                    groups.append(("dims", [d]))
            else:
                groups.append(("dims", dims))
        else:
            groups.append(("dims", dims))
    groups.append(("color", PAGES[3][0]))
    return groups


def build_slides(data, roots):
    """data -> list[slide]；slide = list[元素dict]。"""
    slides = []
    dims_cache = {}
    orig = data["input_image"]
    try:
        with Image.open(orig) as im:
            ow, oh = im.size
    except Exception:
        ow, oh = 800, 600
    osw, osh = fit_box(ow, oh, LEFT_W - 0.1, CONTENT_H)
    ox = LEFT_X + ((LEFT_W - 0.1) - osw) / 2
    oy = CONTENT_Y + (CONTENT_H - osh) / 2

    for sol in data["solutions"]:
        base_title = f"方案{sol['index']} · {sol['direction']}"
        orig_items = [
            {"t": "image", "x": ox, "y": oy, "w": osw, "h": osh, "path": orig},
            {"t": "text", "x": LEFT_X, "y": CONTENT_Y + CONTENT_H - 0.24,
             "w": LEFT_W, "h": 0.22,
             "text": f"测试原图：{data['name']}", "size": 8, "bold": False, "color": GRAY},
        ]
        groups = _page_groups(sol, RIGHT_W)
        n_pages = len(groups)
        for p_idx, (kind, payload) in enumerate(groups, 1):
            if kind == "color":
                # 最后一页：上方 Color 维度（图片放大、文字下移），下方完整整改建议
                dim = payload
                items = [{"t": "text", "x": MARGIN, "y": TITLE_Y, "w": PAGE_W - 2 * MARGIN,
                          "h": TITLE_H, "size": 15, "bold": True, "color": BLACK,
                          "text": f"{base_title}    第{p_idx}/{n_pages}页 · {CN[dim]} + 完整整改建议"}]
                items += orig_items
                items += _dim_block_items(RIGHT_X, CONTENT_Y, RIGHT_W, PAGE4_TOP_H,
                                          sol["ev"].get(dim) or [], LABEL[dim], sol, dim,
                                          roots, dims_cache)
                ay = CONTENT_Y + PAGE4_TOP_H + BLOCK_GAP
                ah = CONTENT_H - PAGE4_TOP_H - BLOCK_GAP
                adv_text = "【该方案完整 · 针对性整改建议】\n" + "\n".join(
                    f"- {CN[d]}整改：{sol['advice'].get(d, '')}" for d in DIM_ORDER)
                pt = fit_pt(adv_text, RIGHT_W, ah, base=10.0, min_pt=7.0)
                items.append({"t": "rect", "x": RIGHT_X - 0.06, "y": ay - 0.06,
                              "w": RIGHT_W + 0.12, "h": ah + 0.12,
                              "fill": (0.965, 0.985, 1.0), "line": (0.55, 0.72, 0.85)})
                items.append({"t": "text", "x": RIGHT_X, "y": ay, "w": RIGHT_W, "h": ah,
                              "text": adv_text, "size": pt, "bold": False, "color": BLACK})
                slides.append(items)
                continue
            dims = payload
            sub = f"{CN[dims[0]]} + {CN[dims[1]]}" if len(dims) == 2 else CN[dims[0]]
            items = [{"t": "text", "x": MARGIN, "y": TITLE_Y, "w": PAGE_W - 2 * MARGIN,
                      "h": TITLE_H, "size": 15, "bold": True, "color": BLACK,
                      "text": f"{base_title}    第{p_idx}/{n_pages}页 · {sub}"}]
            items += orig_items
            free = CONTENT_H - BLOCK_GAP * (len(dims) - 1)
            nat = [_dim_natural(sol["ev"].get(d) or [], sol, d, RIGHT_W) for d in dims]
            # 始终按比例填满整页，让图片尽量占大
            hs = [free * n / sum(nat) for n in nat]
            by = CONTENT_Y
            for dim, bh in zip(dims, hs):
                items += _dim_block_items(RIGHT_X, by, RIGHT_W, bh,
                                          sol["ev"].get(dim) or [], LABEL[dim],
                                          sol, dim, roots, dims_cache)
                by += bh + BLOCK_GAP
            slides.append(items)
    return slides


# ---------------------------------------------------------------------------
# PPTX 渲染
# ---------------------------------------------------------------------------
def _set_font(run, name, size, bold, color):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(int(color[0] * 255), int(color[1] * 255), int(color[2] * 255))
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def render_pptx(slides_items, out_path, font_name="Microsoft YaHei"):
    if not _HAS_PPTX:
        raise RuntimeError("缺少 python-pptx，无法生成 pptx")
    prs = Presentation()
    prs.slide_width = Inches(PAGE_W)
    prs.slide_height = Inches(PAGE_H)
    blank = prs.slide_layouts[6]
    for items in slides_items:
        slide = prs.slides.add_slide(blank)
        for it in items:
            x, y, w, h = (Inches(it["x"]), Inches(it["y"]),
                          Inches(it["w"]), Inches(it["h"]))
            if it["t"] == "image":
                p = it["path"]
                if p and os.path.exists(p):
                    try:
                        slide.shapes.add_picture(p, x, y, w, h)
                        continue
                    except Exception:
                        pass
                sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor(0xE6, 0xE6, 0xE6)
                sh.line.fill.background()
                tf = sh.text_frame
                tf.word_wrap = True
                p_ = tf.paragraphs[0]
                r_ = p_.add_run()
                r_.text = "图片缺失"
                _set_font(r_, font_name, 9, False, (0.3, 0.3, 0.3))
            elif it["t"] == "rect":
                sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor(int(it["fill"][0] * 255),
                                                  int(it["fill"][1] * 255),
                                                  int(it["fill"][2] * 255))
                if it.get("line"):
                    sh.line.color.rgb = RGBColor(int(it["line"][0] * 255),
                                                 int(it["line"][1] * 255),
                                                 int(it["line"][2] * 255))
                    sh.line.width = Pt(0.75)
                else:
                    sh.line.fill.background()
                sh.shadow.inherit = False
            elif it["t"] == "text":
                tb = slide.shapes.add_textbox(x, y, w, h)
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                lines = wrap_text(it["text"], it["w"], it["size"])
                for i, ln in enumerate(lines):
                    p_ = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p_.alignment = PP_ALIGN.LEFT
                    r_ = p_.add_run()
                    r_.text = ln
                    _set_font(r_, font_name, it["size"], it.get("bold", False),
                              it.get("color", BLACK))
    prs.save(out_path)


# ---------------------------------------------------------------------------
# PDF 渲染（reportlab 直渲，无需 LibreOffice）
# ---------------------------------------------------------------------------
def render_pdf(slides_items, out_path):
    if not _HAS_RL:
        raise RuntimeError("缺少 reportlab，无法生成 PDF")
    c = _rl_canvas.Canvas(out_path, pagesize=(PAGE_W * inch, PAGE_H * inch))
    _img_cache = {}

    def _ir(path):
        if path not in _img_cache:
            _img_cache[path] = ImageReader(path)
        return _img_cache[path]

    for items in slides_items:
        for it in items:
            if it["t"] == "image":
                p = it["path"]
                px, py = it["x"] * inch, (PAGE_H - it["y"] - it["h"]) * inch
                pw, ph = it["w"] * inch, it["h"] * inch
                if p and os.path.exists(p):
                    try:
                        c.drawImage(_ir(p), px, py, pw, ph, mask="auto")
                        continue
                    except Exception:
                        pass
                c.setFillColorRGB(0.9, 0.9, 0.9)
                c.rect(px, py, pw, ph, fill=1, stroke=0)
                c.setFillColorRGB(0.35, 0.35, 0.35)
                c.setFont("STSong-Light", 8)
                c.drawCentredString(px + pw / 2, py + ph / 2, "图片缺失")
            elif it["t"] == "rect":
                px, py = it["x"] * inch, (PAGE_H - it["y"] - it["h"]) * inch
                pw, ph = it["w"] * inch, it["h"] * inch
                c.setFillColorRGB(*it["fill"])
                if it.get("line"):
                    c.setStrokeColorRGB(*it["line"])
                    c.setLineWidth(0.75)
                    c.rect(px, py, pw, ph, fill=1, stroke=1)
                else:
                    c.rect(px, py, pw, ph, fill=1, stroke=0)
            elif it["t"] == "text":
                size = it["size"]
                lines = wrap_text(it["text"], it["w"], size)
                x0 = it["x"] * inch
                y = (PAGE_H - it["y"]) * inch - size
                leading = size * 1.25
                c.setFont("STSong-Light", size)
                c.setFillColorRGB(*it.get("color", BLACK))
                for ln in lines:
                    c.drawString(x0, y, ln)
                    y -= leading
        c.showPage()
    c.save()


# ---------------------------------------------------------------------------
# 数据装载
# ---------------------------------------------------------------------------
def load_solution_data(folder, roots) -> dict:
    js_path = os.path.join(folder, "solutions.json")
    with open(js_path, "r", encoding="utf-8") as f:
        d = json.load(f)

    name = os.path.basename(os.path.normpath(folder))
    ev_by_dim = {}
    for dim in DIM_ORDER:
        evs = d.get("evidence_by_dim", {}).get(dim, []) or []
        if EVIDENCE_PER_DIM > 0:
            evs = evs[:EVIDENCE_PER_DIM]
        ev_by_dim[dim] = [{
            "poor": resolve_image(e.get("poor_image_path", ""), roots),
            "good": resolve_image(e.get("good_image_path", ""), roots),
            "text": e.get("dim_original_text", ""),
        } for e in evs]

    solutions = []
    for s in d.get("solutions", []):
        parsed = parse_solution_text(s.get("text", ""))
        solutions.append({
            "index": s.get("index", 0),
            "direction": s.get("direction", ""),
            "defect": parsed["defect"],
            "advice": parsed["advice"],
            "ev": ev_by_dim,
        })
    return {
        "name": name,
        "input_image": resolve_image(d.get("input_image", ""), roots),
        "solutions": solutions,
    }


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------
def _thumb_path(path, thumb_dir, max_side=480):
    """图片降采样到 max_side 内（同一原图只生成一次），用于控制 pptx/pdf 体积。"""
    if not path or not os.path.exists(path):
        return path
    h = hashlib.md5(path.encode("utf-8")).hexdigest()[:16]
    out = os.path.join(thumb_dir, h + ".jpg")
    if not os.path.exists(out):
        try:
            with Image.open(path) as im:
                im = im.convert("RGB")
                im.thumbnail((max_side, max_side))
                im.save(out, "JPEG", quality=82)
        except Exception:
            return path
    return out


def _prepare_thumbs(slides_items):
    """把所有图片路径换成降采样缩略图，返回临时缩略图目录（用后由调用方清理）。"""
    thumb_dir = tempfile.mkdtemp(prefix="ddge_thumbs_")
    for items in slides_items:
        for it in items:
            if it["t"] == "image" and it.get("path"):
                it["path"] = _thumb_path(it["path"], thumb_dir)
    return thumb_dir


def main():
    ap = argparse.ArgumentParser(description="7dimensions 测试结果 -> 单份 PPT + PDF（含全部证据）")
    ap.add_argument("--results-dir", required=True, help="测试结果目录（含各图片名文件夹/solutions.json）")
    ap.add_argument("--save-dir", default=None, help="保存目录（默认=--results-dir，即输入的测试路径）")
    ap.add_argument("--data-root", action="append", default=[],
                    help="数据/图片根目录（可多次指定，用于解析 poor/good/原图缺失路径）")
    ap.add_argument("--only", default=None, help="只处理含该关键字的图片文件夹")
    ap.add_argument("--max-solutions", type=int, default=0, help="每张只取前 N 个方案（0=全部）")
    ap.add_argument("--out-name", default=None,
                    help="输出文件名（不含扩展名，默认=<结果目录名>_报告）")
    ap.add_argument("--skip-pptx", action="store_true", help="不生成 pptx")
    ap.add_argument("--skip-pdf", action="store_true", help="不生成 pdf")
    args = ap.parse_args()

    roots = list(args.data_root) or list(DEFAULT_DATA_ROOTS)
    save_dir = args.save_dir or args.results_dir
    os.makedirs(save_dir, exist_ok=True)

    folders = sorted(
        d for d in glob.glob(os.path.join(args.results_dir, "*", ""))
        if os.path.isfile(os.path.join(d, "solutions.json")))
    if args.only:
        folders = [d for d in folders if args.only in os.path.basename(os.path.normpath(d))]

    if not folders:
        print(f"[make] 在 {args.results_dir} 下没找到含 solutions.json 的文件夹")
        return

    print(f"[make] 待处理图片文件夹: {len(folders)} 个 -> 生成单份报告到 {save_dir}")
    all_slides = []
    for i, folder in enumerate(folders, 1):
        name = os.path.basename(os.path.normpath(folder))
        data = load_solution_data(folder, roots)
        if args.max_solutions > 0:
            data["solutions"] = data["solutions"][:args.max_solutions]
        slides = build_slides(data, roots)
        all_slides += slides
        print(f"[make] ({i}/{len(folders)}) {name}: {len(slides)} 页")

    out_base = os.path.join(
        save_dir,
        args.out_name or (os.path.basename(os.path.normpath(args.results_dir)) + "_报告"))
    print(f"[make] 共 {len(all_slides)} 页，开始渲染 ...")
    thumb_dir = _prepare_thumbs(all_slides)
    try:
        if not args.skip_pptx:
            render_pptx(all_slides, out_base + ".pptx")
        if not args.skip_pdf:
            render_pdf(all_slides, out_base + ".pdf")
    finally:
        shutil.rmtree(thumb_dir, ignore_errors=True)
    print(f"[make] 完成：{len(folders)} 张图片 -> {out_base}.pptx / .pdf（共 {len(all_slides)} 页）")


if __name__ == "__main__":
    main()
